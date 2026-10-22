"""
Telegram 點餐機器人（完整版）
功能：底部常駐按鈕、購物車、飲料可選加珍珠 / 甜度 / 冰塊、結帳、
      訂單通知商家、商家按鈕更新訂單狀態（接單 / 出單 / 完成）並通知客人、
      訂單存檔、在商家群組輸入「結算」產生每日 Excel 報表、
      頻道多篇廣告分時段自動發送（可附圖片）

安裝：py -m pip install "python-telegram-bot[job-queue]>=20" tzdata openpyxl pillow
執行：py order_bot.py
"""
import asyncio
import csv
import io
import json
import logging
import math
import os
import re
import socket
import time as _time
import zipfile
from datetime import datetime, time, timedelta, timezone
from zoneinfo import ZoneInfo

from telegram import (
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    InputMediaPhoto,
    ReplyKeyboardMarkup,
    ReplyKeyboardRemove,
    Update,
)
from telegram.error import BadRequest, Conflict, NetworkError, RetryAfter
from telegram.ext import (
    Application,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    ConversationHandler,
    MessageHandler,
    filters,
)

logging.basicConfig(level=logging.INFO)
logging.getLogger("httpx").setLevel(logging.WARNING)  # 避免日誌印出 token

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
# 會變動的資料（訂單紀錄、設定、庫存、報表）放在 DATA_DIR。
# 雲端平台請把環境變數 DATA_DIR 設成「永久儲存空間」的掛載路徑（例如 /data），重新部署資料才不會消失。
DATA_DIR = os.getenv("DATA_DIR") or BASE_DIR
os.makedirs(DATA_DIR, exist_ok=True)
VERSION = "2026-10-22 v17（關店後菜單一律顯示已休店）"

# ==================== 設定區（請修改）====================
TOKEN = os.getenv("BOT_TOKEN", "在這裡貼上BOT_TOKEN")
ADMIN_CHAT_ID = int(os.getenv("ADMIN_CHAT_ID", "-1004424236242"))  # 接收訂單的商家個人或群組 id

# 以下幾項也可以用同名的環境變數設定（雲端平台不用改程式）；SHOP_CONTACT 的換行寫成 \n
SHOP_NAME = os.getenv("SHOP_NAME", "雞排王")
SHOP_CONTACT = os.getenv("SHOP_CONTACT", "客服：@chicken5757\n營業時間：18:00 - 03:00").replace("\\n", "\n")

# 客服的 Telegram 使用者名稱（不含 @，只能英文數字底線）；留空 "" 就不顯示「聯繫客服」按鈕
CS_USERNAME = os.getenv("CS_USERNAME", "")

# ---------- 頻道廣告（CHANNEL_ID 留空 "" 就不啟用）----------
CHANNEL_ID = os.getenv("CHANNEL_ID", "")  # 例如 "@my_channel" 或 "-100xxxxxxxxxx"
BOT_USERNAME = os.getenv("BOT_USERNAME", "你的機器人使用者名稱")  # 不含 @
TIMEZONE = ZoneInfo("Asia/Kuala_Lumpur")
# 「營業日」從幾點開始算新的一天（要落在收攤之後、開店之前）。
# 營業時間 18:00~03:00 → 設 17：晚上 18:00 到隔天 03:00 的訂單都算同一天，
# 而且 03:00 之後、17:00 之前輸入「結算」，統計的都是昨晚這一班。
DAY_START_HOUR = 17
TEST_AD_ON_START = False  # 改成 True：啟動 10 秒後先發第 1 篇廣告測試（測完改回 False）

# 每一篇廣告：times 發送時間（可多個）、photo 圖片檔名（放同一資料夾，不要圖片填 ""）、text 文案
ADS = [
    {   # 第 1 篇：中午 11:00（雞排）
        "times": [time(11, 0)],
        "photo": "ad1.jpg",
        "text": (
            "🇹🇼 人在外地久了，最想念的，往往就是這一口台灣味 🇹🇼\n\n"
            "還記得以前放學、下班後，經過夜市聞到炸雞香味\n\n"
            "🍗外皮炸得金黃酥脆 裡面的雞肉嫩嫩多汁🍗\n\n"
            "🍗 現點現炸\n"
            "🔥 外酥內嫩\n"
            "❤️ 一口找回熟悉的家鄉味\n\n"
            "🔥 雞排王｜一口咬下，滿滿台灣味！\n\n"
            "#雞排王 #台灣味 #家鄉味 #台灣炸雞"
        ),
    },
    {   # 第 2 篇：傍晚 17:00（飲料）
        "times": [time(17, 0)],
        "photo": "ad2.jpg",
        "text": (
            "🇹🇼 想念台灣，就喝這一杯 🧋\n\n"
            "人在異鄉久了🌏最想念的不是什麼山珍海味，\n"
            "而是那熟悉的 台灣茶香與家鄉味 🍵🏠❤️\n\n"
            "一杯阿薩姆紅茶 🫖\n"
            "一杯高山烏龍 🌿\n"
            "再來上一杯珍珠奶茶的熟悉味道 🧋✨\n\n"
            "不用等回台灣 🇹🇼 想念的味道，現在就喝得到 ❤️\n\n"
            "一口茶香，一口家鄉味 🏠🇹🇼"
        ),
    },
]

CURRENCY = "U"
AUTO_BACKUP_TIME = time(4, 0)  # 每天這個時間，把資料備份檔傳到商家群組；None = 不自動備份（關店時一定會傳）
RM_ROUND_5_SEN = True  # 「合計」的馬幣金額四捨五入到 5 仙（馬來西亞現金找零習慣）；單品只四捨五入到 1 仙

# 菜單：分類 -> {代號: (名稱, 價格)}
MENU = {
    "🍗 肉類": {
        "m1": ("雞翅", 1.45),
        "m2": ("雞腿", 2.45),
        "m3": ("腿排", 2.45),
        "m4": ("古早味雞排", 4.45),
        "m5": ("塔香鹽酥雞", 3.20),
    },
    "🍡 點心": {
        "s1": ("甜不辣", 2.00),
        "s2": ("薯條", 2.00),
        "s3": ("杏鮑菇", 2.20),
    },
    "🥤 飲料": {
        "d1": ("阿薩姆紅茶", 1.95),
        "d2": ("甘露綠茶", 1.95),
        "d3": ("高山烏龍茶", 1.95),
        "d4": ("阿薩姆奶茶", 2.45),
        "d5": ("甘露奶綠", 2.45),
        "d6": ("高山烏龍奶", 2.45),
        "d7": ("珍珠奶茶", 2.70),
        "d8": ("巷口冬瓜露", 1.95),
    },
}

# 飲料分類：點選後依序詢問 加珍珠 → 甜度 → 冰塊
DRINK_CATEGORY = "🥤 飲料"
TOPPING_NAME = "加珍珠"
TOPPING_PRICE = 0.70
NO_TOPPING = {"d7"}  # 不詢問配料的品項（珍珠奶茶本來就有珍珠）
SUGAR = ["1分", "3分", "5分", "7分", "10分"]
ICE = ["去冰", "微冰", "少冰", "正常冰"]
# 甜度固定的飲料：客人不用選甜度，只選冰塊。
# 左邊寫「品名」（要跟上面菜單的品名一模一樣），右邊是顯示在購物車和杯貼上的文字（寫 "" 就不顯示）
FIXED_SUGAR_ITEMS = {
    "巷口冬瓜露": "固定甜度",
    "冬瓜茶": "固定甜度",
}

# ---------- 杯貼（NIIMBOT B21，貼紙 50×30mm）----------
STICKER_ENABLED = True      # 商家按「接單」時，自動把飲料貼紙傳到商家群組
STICKER_SEND_IMAGES = True  # 傳「一杯一張」的貼紙圖片
STICKER_SEND_EXCEL = True   # 另傳一個 Excel（一杯一列），給 NIIMBOT App 的 Excel 批量列印用
STICKER_ALL_ITEMS = False   # True：連餐點（雞排等）也做貼紙；False：只做飲料
STICKER_SHOW_NAME = True    # 貼紙上是否印客人姓名
STICKER_SCALE = 2           # 圖片倍率：2 = 800×480 像素（建議）；1 = 400×240 像素
STICKER_FONT = ""           # 中文字型檔路徑；留空 = 自動尋找（Windows 會用微軟正黑體）
# ==========================================================

LOG_FILE = os.path.join(DATA_DIR, "orders_log.jsonl")  # 每筆訂單與狀態的完整紀錄（結算報表的資料來源）
REPORT_DIR = os.path.join(DATA_DIR, "reports")  # 結算報表存放資料夾
ORDER_LOCK = asyncio.Lock()
START_TIME = datetime.now(TIMEZONE)
ERROR_SEEN = {}  # 錯誤提醒的節流（避免一直洗版）
ADMIN_FILTER = None  # 商家群組的過濾器，在 main() 建立（群組升級時會動態更新）


def run_info() -> str:
    """這份機器人跑在哪一台電腦上（用來分辨是不是有人同時開了好幾份）"""
    return (
        f"🖥 執行電腦：{socket.gethostname()}（PID {os.getpid()}）\n"
        f"⏱ 這一份啟動於：{START_TIME.strftime('%m-%d %H:%M:%S')}\n"
        f"📁 資料位置：{DATA_DIR}"
    )

# 外帶 / 外送開關（在商家群組輸入「關閉外帶」「開啟外帶」切換，重啟後仍保留）
SETTINGS_FILE = os.path.join(DATA_DIR, "settings.json")


DEFAULT_SETTINGS = {
    "pickup": True,      # 外帶是否開放
    "delivery": True,    # 外送是否開放
    "open": True,        # 是否營業中（關店後客人不能下單）
    "session": None,     # 目前（或最近一次）營業場次的名稱，開店時建立
    "opened_at": None,   # 開店時間
    "closed_at": None,   # 關店時間
    "rm_rate": None,     # 匯率：1U 等於幾馬幣（RM）。None = 不顯示馬幣
    "admin_migrated_from": None,  # 商家群組升級成超級群組時，記錄舊 id
    "admin_migrated_to": None,    # ……與新 id
}


def load_settings() -> dict:
    data = dict(DEFAULT_SETTINGS)
    try:
        with open(SETTINGS_FILE, encoding="utf-8") as f:
            saved = json.load(f)
        for k, default in DEFAULT_SETTINGS.items():
            if k in saved:
                data[k] = bool(saved[k]) if isinstance(default, bool) else saved[k]
    except (OSError, ValueError):
        pass
    return data


SETTINGS = load_settings()
# 商家群組升級成超級群組後 id 會改變：如果設定檔裡的 id 就是升級前的舊 id，自動改用新 id
if SETTINGS.get("admin_migrated_from") == ADMIN_CHAT_ID and SETTINGS.get("admin_migrated_to"):
    ADMIN_CHAT_ID = int(SETTINGS["admin_migrated_to"])


def save_settings():
    try:
        with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
            json.dump(SETTINGS, f)
    except OSError:
        logging.exception("寫入 settings.json 失敗")

# ---------- 每日庫存 ----------
# stock.json 存「每日可賣數量」(limits) 與「手動標示售完」(manual)。
# 已賣數量不另外存，直接由訂單紀錄計算（拒單會自動把庫存還回去）。
STOCK_FILE = os.path.join(DATA_DIR, "stock.json")
LOW_STOCK_HINT = 5  # 剩下幾份以內，菜單上會顯示「剩N」


def load_stock() -> dict:
    data = {"limits": {}, "manual": {"label": None, "items": []}}
    try:
        with open(STOCK_FILE, encoding="utf-8") as f:
            saved = json.load(f)
        data["limits"] = {k: int(v) for k, v in saved.get("limits", {}).items()}
        m = saved.get("manual", {})
        data["manual"] = {"label": m.get("label"), "items": list(m.get("items", []))}
    except (OSError, ValueError, TypeError, AttributeError):
        pass
    return data


STOCK = load_stock()


def save_stock():
    try:
        with open(STOCK_FILE, "w", encoding="utf-8") as f:
            json.dump(STOCK, f, ensure_ascii=False)
    except OSError:
        logging.exception("寫入 stock.json 失敗")


# 建立品項索引：加珍珠的版本代號 = 原代號 + "p"
ITEMS = {}
ITEM_CAT = {}
for _cat, _group in MENU.items():
    for _iid, (_name, _price) in _group.items():
        ITEMS[_iid] = (_name, _price)
        ITEM_CAT[_iid] = _cat
        if _cat == DRINK_CATEGORY and _iid not in NO_TOPPING:
            ITEMS[_iid + "p"] = (f"{_name}({TOPPING_NAME})", round(_price + TOPPING_PRICE, 2))
            ITEM_CAT[_iid + "p"] = _cat


FIXED_SUGAR = {  # 品項代號 → 顯示文字
    _iid: FIXED_SUGAR_ITEMS[_name]
    for _group in MENU.values()
    for _iid, (_name, _price) in _group.items()
    if _name in FIXED_SUGAR_ITEMS
}


def is_drink(iid: str) -> bool:
    return ITEM_CAT.get(iid) == DRINK_CATEGORY


def base_id(iid: str) -> str:
    return iid[:-1] if iid.endswith("p") else iid


def asks_topping(iid: str) -> bool:
    return is_drink(iid) and iid not in NO_TOPPING and not iid.endswith("p")


# 購物車的鍵：一般品項 "m1"；飲料 "d1p.2.3"（品項代號.甜度序號.冰塊序號）
def parse_key(key: str):
    if "." in key:
        iid, s, i = key.split(".")
        return iid, int(s), int(i)
    return key, None, None


def item_price(key: str) -> float:
    return ITEMS[parse_key(key)[0]][1]


def is_fixed_sugar(iid: str) -> bool:
    return base_id(iid) in FIXED_SUGAR


def sugar_raw(iid: str, s) -> str:
    """甜度的原始文字：選的是「5分」；甜度固定的飲料（s = -1）是設定的固定文字"""
    if s is None:
        return ""
    if s < 0:
        return FIXED_SUGAR.get(base_id(iid), "")
    return SUGAR[s]


def sugar_text(raw: str) -> str:
    """顯示用：「5分」→「5分糖」；固定甜度的文字照原樣"""
    return f"{raw}糖" if raw in SUGAR else raw


def display_name(key: str) -> str:
    iid, s, i = parse_key(key)
    name = ITEMS[iid][0]
    if s is not None:
        parts = [p for p in (sugar_text(sugar_raw(iid, s)), ICE[i]) if p]
        name += f"（{'・'.join(parts)}）"
    return name


# 底部常駐按鈕
BTN_MENU = "🍽 菜單"
BTN_CART = "🛒 購物車"
BTN_CHECKOUT = "✅ 結帳"
BTN_ORDERS = "📋 我的訂單"
BTN_CONTACT = "☎️ 聯絡店家"
BTN_CANCEL = "❌ 取消"
BTN_BACK_MENU = "🍽 返回菜單"
BTN_PICKUP = "🏃 外帶"
BTN_DELIVERY = "🛵 外送"
BTN_CONFIRM = "✅ 確認送出"

MAIN_KB = ReplyKeyboardMarkup(
    [[BTN_MENU, BTN_CART], [BTN_CHECKOUT, BTN_ORDERS], [BTN_CONTACT]],
    resize_keyboard=True,
    is_persistent=True,
)

# 結帳流程狀態
TYPE, NAME, PHONE, ADDR, CONFIRM = range(5)


# ==================== 工具函式 ====================
def rm_amount(u: float, rate, total: bool = False):
    """把 U 換成馬幣。rate = 1U 等於幾 RM；沒設定匯率回傳 None"""
    if not rate:
        return None
    sen = int(math.floor(u * rate * 100 + 0.5))  # 先四捨五入到 1 仙
    if total and RM_ROUND_5_SEN:  # 再依馬來西亞規則進位到 5 仙
        r = sen % 5
        if r in (1, 2):
            sen -= r
        elif r in (3, 4):
            sen += 5 - r
    return sen / 100


def money(x: float, total: bool = False) -> str:
    """金額文字：有設定匯率時，例如 4.45U（RM19.80）"""
    base = f"{x:.2f}{CURRENCY}"
    rm = rm_amount(x, SETTINGS.get("rm_rate"), total)
    return base if rm is None else f"{base}（RM{rm:.2f}）"


def get_cart(ctx: ContextTypes.DEFAULT_TYPE) -> dict:
    return ctx.user_data.setdefault("cart", {})


def cart_total(cart: dict) -> float:
    return round(sum(item_price(k) * q for k, q in cart.items()), 2)


def cart_text(cart: dict) -> str:
    lines = [
        f"• {display_name(k)} x{q}  {money(item_price(k) * q)}" for k, q in cart.items()
    ]
    return "\n".join(lines) + f"\n\n💰 合計：{money(cart_total(cart), total=True)}"


def render_cart(ctx):
    cart = get_cart(ctx)
    if not cart:
        return "🛒 購物車是空的，請先選餐。", InlineKeyboardMarkup(
            [[InlineKeyboardButton("🍽 去點餐", callback_data="menu")]]
        )
    rows = []
    for key, q in cart.items():
        rows.append([
            InlineKeyboardButton("➖", callback_data=f"dec:{key}"),
            InlineKeyboardButton(f"{display_name(key)} x{q}", callback_data="noop"),
            InlineKeyboardButton("➕", callback_data=f"inc:{key}"),
        ])
    rows.append([
        InlineKeyboardButton("🗑 清空", callback_data="clear"),
        InlineKeyboardButton("✅ 去結帳", callback_data="checkout"),
    ])
    rows.append([InlineKeyboardButton("🍽 返回菜單", callback_data="menu")])
    return "🛒 你的購物車\n\n" + cart_text(cart), InlineKeyboardMarkup(rows)


def category_markup():
    rows = [[InlineKeyboardButton(c, callback_data=f"cat:{c}")] for c in MENU]
    return InlineKeyboardMarkup(rows)


def category_view(cat):
    rows = []
    for iid, (n, p) in MENU[cat].items():
        r = remaining(iid)
        if r is not None and r <= 0:
            rows.append([InlineKeyboardButton(f"❌ {n}（已售完）", callback_data="sold")])
            continue
        cb = f"pick:{iid}" if is_drink(iid) else f"add:{iid}"
        tail = f"（剩{r}）" if r is not None and r <= LOW_STOCK_HINT else ""
        rows.append([InlineKeyboardButton(f"{n}  {money(p)}{tail}", callback_data=cb)])
    rows.append([
        InlineKeyboardButton("⬅️ 返回分類", callback_data="menu"),
        InlineKeyboardButton("🛒 購物車", callback_data="cart"),
    ])
    return f"{cat}\n點選品項加入購物車：", InlineKeyboardMarkup(rows)


# ----- 飲料選項：加珍珠 → 甜度 → 冰塊 -----
def topping_view(iid):
    name, price = ITEMS[iid]
    with_topping = ITEMS[iid + "p"][1]
    rows = [
        [InlineKeyboardButton(f"不加  {money(price)}", callback_data=f"sg:{iid}")],
        [InlineKeyboardButton(
            f"🧋 {TOPPING_NAME} +{money(TOPPING_PRICE)}  → {money(with_topping)}",
            callback_data=f"sg:{iid}p",
        )],
        [InlineKeyboardButton("⬅️ 返回", callback_data=f"cat:{ITEM_CAT[iid]}")],
    ]
    return f"{name}\n要{TOPPING_NAME}嗎？", InlineKeyboardMarkup(rows)


def sugar_view(iid):
    name = ITEMS[iid][0]
    btns = [
        InlineKeyboardButton(label, callback_data=f"ic:{iid}:{s}")
        for s, label in enumerate(SUGAR)
    ]
    back_iid = base_id(iid)
    back = f"pick:{back_iid}" if asks_topping(back_iid) else f"cat:{ITEM_CAT[iid]}"
    rows = [btns[:3], btns[3:], [InlineKeyboardButton("⬅️ 返回", callback_data=back)]]
    return f"{name}\n請選擇甜度：", InlineKeyboardMarkup(rows)


def after_topping(iid):
    """加珍珠（或不問加珍珠）之後的下一步：甜度固定的飲料直接問冰塊"""
    return ice_view(iid, -1) if is_fixed_sugar(iid) else sugar_view(iid)


def ice_view(iid, s):
    name = ITEMS[iid][0]
    btns = [
        InlineKeyboardButton(label, callback_data=f"add:{iid}.{s}.{i}")
        for i, label in enumerate(ICE)
    ]
    if is_fixed_sugar(iid):  # 沒有甜度這一步，返回就回到加珍珠或品項列表
        back_iid = base_id(iid)
        back = f"pick:{back_iid}" if asks_topping(back_iid) else f"cat:{ITEM_CAT[iid]}"
    else:
        back = f"sg:{iid}"
    rows = [btns[:2], btns[2:], [InlineKeyboardButton("⬅️ 返回", callback_data=back)]]
    sugar_part = sugar_text(sugar_raw(iid, s))
    head = f"{name}（{sugar_part}）" if sugar_part else name
    return f"{head}\n請選擇冰塊：", InlineKeyboardMarkup(rows)


# ==================== 訂單紀錄（結算報表的資料來源）====================
def business_day(dt=None) -> str:
    """時間換算成「營業日」日期：DAY_START_HOUR 點之前的訂單，算前一天"""
    dt = dt or datetime.now(TIMEZONE)
    return (dt - timedelta(hours=DAY_START_HOUR)).strftime("%Y-%m-%d")


def today_str() -> str:
    """目前的場次名稱：用過「開店」就是該場次；還沒用過就沿用依時間換算的營業日"""
    return SETTINGS["session"] or business_day()


def valid_day_label(day: str) -> bool:
    """場次名稱格式：2026-10-09 或 2026-10-09-2"""
    m = re.fullmatch(r"(\d{4}-\d{2}-\d{2})(?:-\d+)?", day)
    if not m:
        return False
    try:
        datetime.strptime(m.group(1), "%Y-%m-%d")
        return True
    except ValueError:
        return False


def log_event(rec: dict):
    """把一筆事件（新訂單 / 狀態變更）附加寫入 orders_log.jsonl"""
    try:
        with open(LOG_FILE, "a", encoding="utf-8") as f:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")
    except Exception:
        logging.exception("寫入 orders_log.jsonl 失敗")


def load_all_orders() -> list:
    """讀取紀錄裡的所有訂單，並套用最新的訂單狀態"""
    if not os.path.exists(LOG_FILE):
        return []
    orders, latest = [], {}
    with open(LOG_FILE, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
            except ValueError:
                continue
            if rec.get("event") == "order":
                rec["status"] = "new"
                latest[(rec["uid"], rec["no"])] = rec
                orders.append(rec)
            elif rec.get("event") == "status":
                o = latest.get((rec.get("uid"), rec.get("no")))
                if o:
                    o["status"] = rec["status"]
    return orders


def order_day(o: dict) -> str:
    """訂單所屬場次：有開店場次就用它；舊訂單有完整下單時間就依 DAY_START_HOUR 換算；更舊的沿用當時寫下的日期"""
    if o.get("session"):  # 開店後下的單，直接用當時的場次
        return o["session"]
    ts = o.get("ts")
    if ts:
        try:
            return business_day(datetime.fromisoformat(ts))
        except ValueError:
            pass
    return o.get("date", "")


def load_orders_for(day: str) -> list:
    return [o for o in load_all_orders() if order_day(o) == day]


def next_order_no() -> int:
    """今天的下一個訂單編號（依紀錄計算，機器人重啟也不會重號）"""
    nums = []
    for o in load_orders_for(today_str()):
        try:
            nums.append(int(o["no"]))
        except (KeyError, ValueError):
            pass
    return max(nums, default=0) + 1


_SOLD_CACHE = {}


def sold_counts(day: str) -> dict:
    """某場次各品項已賣份數（不含拒單）。訂單紀錄沒變動就直接用快取，所以很快也很準"""
    try:
        st = os.stat(LOG_FILE)
        sig = (day, st.st_mtime_ns, st.st_size)
    except OSError:
        sig = (day, None, None)
    if _SOLD_CACHE.get("sig") == sig:
        return _SOLD_CACHE["val"]
    counts = {}
    for o in load_orders_for(day):
        if o.get("status") == "reject":
            continue
        for l in o["lines"]:
            counts[l["iid"]] = counts.get(l["iid"], 0) + l["qty"]
    _SOLD_CACHE.update(sig=sig, val=counts)
    return counts


def manual_soldout() -> set:
    m = STOCK["manual"]
    return set(m["items"]) if m.get("label") == today_str() else set()


def remaining(iid: str):
    """剩餘份數。None 代表不限量；0 代表已售完"""
    if iid in manual_soldout():
        return 0
    limit = STOCK["limits"].get(iid)
    if limit is None:
        return None
    return max(limit - sold_counts(today_str()).get(iid, 0), 0)


def is_sold_out(iid: str) -> bool:
    r = remaining(iid)
    return r is not None and r <= 0


def cart_qty(cart: dict, bid: str) -> int:
    return sum(q for k, q in cart.items() if base_id(parse_key(k)[0]) == bid)


def can_add(cart: dict, bid: str):
    """再加一份會不會超過庫存。回傳 (可以嗎, 說明)"""
    r = remaining(bid)
    if r is None:
        return True, ""
    name = ITEMS[bid][0]
    if r <= 0:
        return False, f"😢 {name} 今天已售完"
    have = cart_qty(cart, bid)
    if have + 1 > r:
        return False, f"{name} 今天只剩 {r} 份，你已經選了 {have} 份"
    return True, ""


def stock_problems(cart: dict) -> list:
    """購物車裡庫存不足的品項（結帳前與送出訂單時檢查）"""
    need = {}
    for k, q in cart.items():
        bid = base_id(parse_key(k)[0])
        need[bid] = need.get(bid, 0) + q
    out = []
    for bid, n in need.items():
        r = remaining(bid)
        if r is None:
            continue
        name = ITEMS[bid][0]
        if r <= 0:
            out.append(f"• {name}：今天已售完")
        elif n > r:
            out.append(f"• {name}：只剩 {r} 份（你選了 {n} 份）")
    return out


def order_lines(cart: dict) -> list:
    out = []
    for key, q in cart.items():
        iid, s, i = parse_key(key)
        bid = base_id(iid)
        out.append({
            "iid": bid,
            "name": ITEMS[bid][0],
            "topping": iid.endswith("p"),
            "sugar": sugar_raw(iid, s),
            "ice": ICE[i] if i is not None else "",
            "qty": q,
            "price": ITEMS[iid][1],
        })
    return out


def save_order(no, d, cart):
    """把訂單寫入 orders.csv（Excel 可直接開啟）"""
    items = "; ".join(f"{display_name(k)}x{q}" for k, q in cart.items())
    path = os.path.join(DATA_DIR, "orders.csv")
    new_file = not os.path.exists(path) or os.path.getsize(path) == 0
    with open(path, "a", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        if new_file:
            w.writerow(["日期時間", "單號", "姓名", "電話", "方式", "地址", "內容", f"金額（{CURRENCY}）"])
        w.writerow([
            datetime.now(TIMEZONE).strftime("%Y-%m-%d %H:%M"),
            f"#{no}", d["name"], d["phone"], d["order_type"], d["addr"],
            items, cart_total(cart),
        ])


# ==================== 基本指令 ====================
MERCHANT_HELP = (
    "🧑‍🍳 商家指令（只有商家群組能用）\n\n"
    "🏪 營業\n開店 ／ 關店\n匯率 4.45（設定 1U 等於幾馬幣）\n結算（查舊場次：結算 2026-10-10）\n訂單（只傳訂單清單）\n營業日（查看目前場次）\n\n"
    "📦 庫存\n庫存\n設定庫存（換行寫「雞排 50」）\n售完 雞排 ／ 補貨 雞排 ／ 清除庫存\n\n"
    "🛵 服務\n關閉外帶 ／ 開啟外帶\n關閉外送 ／ 開啟外送\n\n"
    "💾 資料\n備份（傳資料備份檔，關店也會自動傳）\n匯入：傳備份檔並在說明寫「匯入」\n\n"
    "🏷 杯貼\n接單時自動傳杯貼\n補印 12（重傳 #0012）\n貼紙測試（傳範例貼紙）\n\n"
    "🔧 其他\n/ping 速度與執行電腦\n/chatid 查看群組 id\n隱藏選單：收起點餐按鈕"
)


async def hide_menu(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """商家群組輸入「隱藏選單」或 /hidemenu：把機器人的點餐按鈕從這個群組收起來"""
    await update.effective_message.reply_text(
        "✅ 已收起機器人的點餐按鈕。", reply_markup=ReplyKeyboardRemove()
    )


async def start(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if ADMIN_CHAT_ID and update.effective_chat.id == ADMIN_CHAT_ID:
        # 商家群組：不顯示客人的點餐按鈕，順便把已經出現的收起來
        await update.message.reply_text(MERCHANT_HELP, reply_markup=ReplyKeyboardRemove())
        return
    text = f"歡迎光臨「{SHOP_NAME}」👋\n請使用下方按鈕開始點餐。"
    if not SETTINGS["open"]:
        text += "\n\n🔴 已休店，目前暫停接單，營業時間再來喔！"
    await update.message.reply_text(text, reply_markup=MAIN_KB)


async def show_menu(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if await closed_guard(update):
        return
    await update.message.reply_text("請選擇分類：", reply_markup=category_markup())


async def show_cart(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if await closed_guard(update):
        return
    text, markup = render_cart(ctx)
    await update.message.reply_text(text, reply_markup=markup)


async def show_contact(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(f"☎️ {SHOP_NAME}\n{SHOP_CONTACT}")


async def show_orders(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    orders = ctx.user_data.get("orders", [])
    if not orders:
        await update.message.reply_text("目前沒有訂單紀錄。")
        return
    await update.message.reply_text("📋 最近訂單\n\n" + "\n\n".join(orders[-5:]))


# ==================== 客人端：內嵌按鈕 ====================
async def on_callback(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    data = q.data
    cart = get_cart(ctx)

    if not SETTINGS["open"] and data != "noop":
        # 已休店：不管按哪個按鈕，都跳出「已休店」，並把這則菜單訊息換成休店說明（按鈕消失）
        await q.answer("🔴 已休店，目前暫停接單", show_alert=True)
        try:
            await q.edit_message_text(closed_text())
        except BadRequest:
            pass
        return

    if data == "noop":
        await q.answer()

    elif data == "menu":
        await q.answer()
        await q.edit_message_text("請選擇分類：", reply_markup=category_markup())

    elif data.startswith("cat:"):
        await q.answer()
        text, markup = category_view(data[4:])
        await q.edit_message_text(text, reply_markup=markup)

    elif data == "sold":
        await q.answer("😢 這項今天已售完，請選擇其他品項", show_alert=True)

    elif data.startswith("pick:"):
        # 飲料：先問加珍珠（有的話），再問甜度
        iid = data[5:]
        ok, why = can_add(cart, base_id(iid))
        if not ok:
            await q.answer(why, show_alert=True)
            try:
                text, markup = category_view(ITEM_CAT[iid])
                await q.edit_message_text(text, reply_markup=markup)
            except BadRequest:
                pass
            return
        text, markup = topping_view(iid) if asks_topping(iid) else after_topping(iid)
        await q.answer()
        await q.edit_message_text(text, reply_markup=markup)

    elif data.startswith("sg:"):
        await q.answer()
        text, markup = after_topping(data[3:])
        await q.edit_message_text(text, reply_markup=markup)

    elif data.startswith("ic:"):
        _, iid, s = data.split(":")
        await q.answer()
        text, markup = ice_view(iid, int(s))
        await q.edit_message_text(text, reply_markup=markup)

    elif data.startswith("add:"):
        key = data[4:]
        bid = base_id(parse_key(key)[0])
        ok, why = can_add(cart, bid)
        if not ok:
            await q.answer(why, show_alert=True)
            try:
                text, markup = category_view(ITEM_CAT[bid])
                await q.edit_message_text(text, reply_markup=markup)
            except BadRequest:
                pass
            return
        cart[key] = cart.get(key, 0) + 1
        await q.answer(f"已加入：{display_name(key)}（共 {cart[key]} 份）", show_alert=False)
        # 飲料選完冰塊後，回到飲料列表方便繼續點
        if "." in key:
            try:
                text, markup = category_view(DRINK_CATEGORY)
                await q.edit_message_text(text, reply_markup=markup)
            except BadRequest:
                pass

    elif data.startswith("inc:") or data.startswith("dec:"):
        action, key = data.split(":")
        if action == "inc":
            ok, why = can_add(cart, base_id(parse_key(key)[0]))
            if not ok:
                await q.answer(why, show_alert=True)
                return
            cart[key] = cart.get(key, 0) + 1
        else:
            cart[key] = cart.get(key, 0) - 1
            if cart[key] <= 0:
                cart.pop(key, None)
        await q.answer()
        text, markup = render_cart(ctx)
        await q.edit_message_text(text, reply_markup=markup)

    elif data == "clear":
        cart.clear()
        await q.answer("已清空")
        text, markup = render_cart(ctx)
        await q.edit_message_text(text, reply_markup=markup)

    elif data == "cart":
        await q.answer()
        text, markup = render_cart(ctx)
        await q.edit_message_text(text, reply_markup=markup)


# ==================== 結帳流程 ====================
async def checkout_start(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if update.callback_query:
        await update.callback_query.answer()
    msg = update.effective_message
    if not get_cart(ctx):
        await msg.reply_text("購物車是空的，請先點「🍽 菜單」選餐。", reply_markup=MAIN_KB)
        return ConversationHandler.END
    if not SETTINGS["open"]:
        await msg.reply_text(closed_text(), reply_markup=MAIN_KB)
        return ConversationHandler.END
    problems = stock_problems(get_cart(ctx))
    if problems:
        await msg.reply_text(stock_warning(problems), reply_markup=MAIN_KB)
        return ConversationHandler.END
    options = []
    if SETTINGS["pickup"]:
        options.append(BTN_PICKUP)
    if SETTINGS["delivery"]:
        options.append(BTN_DELIVERY)
    if not options:
        await msg.reply_text(
            "😢 目前暫時停止接單，請稍後再來，或聯絡店家：\n" + SHOP_CONTACT,
            reply_markup=MAIN_KB,
        )
        return ConversationHandler.END
    notice = ""
    if not SETTINGS["pickup"]:
        notice = "（目前暫停外帶，僅開放外送）\n"
    elif not SETTINGS["delivery"]:
        notice = "（目前暫停外送，僅開放外帶）\n"
    kb = ReplyKeyboardMarkup([options, [BTN_BACK_MENU, BTN_CANCEL]], resize_keyboard=True)
    await msg.reply_text(f"{notice}請選擇取餐方式：", reply_markup=kb)
    return TYPE


async def got_type(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if await closed_guard(update):
        return ConversationHandler.END
    choice = update.message.text
    if (choice == BTN_PICKUP and not SETTINGS["pickup"]) or (
        choice == BTN_DELIVERY and not SETTINGS["delivery"]
    ):
        await update.message.reply_text("這個取餐方式目前暫停服務，請選擇其他方式。")
        return TYPE
    ctx.user_data["order_type"] = choice
    await update.message.reply_text(
        "請輸入你的姓名：",
        reply_markup=ReplyKeyboardMarkup([[BTN_BACK_MENU, BTN_CANCEL]], resize_keyboard=True),
    )
    return NAME


async def got_name(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if await closed_guard(update):
        return ConversationHandler.END
    ctx.user_data["name"] = update.message.text.strip()
    await update.message.reply_text("請輸入聯絡電話：")
    return PHONE


async def got_phone(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if await closed_guard(update):
        return ConversationHandler.END
    ctx.user_data["phone"] = update.message.text.strip()
    if ctx.user_data["order_type"] == BTN_DELIVERY:
        await update.message.reply_text("請輸入外送地址：")
        return ADDR
    ctx.user_data["addr"] = "—（外帶）"
    return await ask_confirm(update, ctx)


async def got_addr(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if await closed_guard(update):
        return ConversationHandler.END
    ctx.user_data["addr"] = update.message.text.strip()
    return await ask_confirm(update, ctx)


async def ask_confirm(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    d = ctx.user_data
    text = (
        "📝 請確認訂單\n\n"
        f"{cart_text(get_cart(ctx))}\n\n"
        f"方式：{d['order_type']}\n姓名：{d['name']}\n電話：{d['phone']}\n地址：{d['addr']}"
    )
    kb = ReplyKeyboardMarkup([[BTN_CONFIRM], [BTN_BACK_MENU, BTN_CANCEL]], resize_keyboard=True)
    await update.message.reply_text(text, reply_markup=kb)
    return CONFIRM


def stock_warning(problems: list) -> str:
    return (
        "⚠️ 庫存有變動，請先調整購物車：\n" + "\n".join(problems)
        + "\n\n請按下方「🛒 購物車」調整數量，或「🍽 菜單」選其他品項。"
    )


def closed_text() -> str:
    return "🔴 已休店\n目前暫停接單，歡迎營業時間再來點餐！\n\n" + SHOP_CONTACT


async def closed_guard(update: Update) -> bool:
    """已休店時回覆「已休店」並回傳 True，讓呼叫的地方直接結束"""
    if SETTINGS["open"]:
        return False
    await update.effective_message.reply_text(closed_text(), reply_markup=MAIN_KB)
    return True


async def send_to_admin(bot, text, markup) -> bool:
    """把訂單傳給商家。遇到 Telegram 限流（太多訊息）會等待後重試，最多 5 次"""
    for _ in range(5):
        try:
            await bot.send_message(ADMIN_CHAT_ID, text, reply_markup=markup)
            return True
        except RetryAfter as e:
            wait = e.retry_after
            wait = wait.total_seconds() if hasattr(wait, "total_seconds") else float(wait)
            logging.warning("商家群組訊息太頻繁，%.0f 秒後重試", wait)
            await asyncio.sleep(wait + 1)
        except Exception:
            logging.exception("無法把訂單傳給商家，請檢查 ADMIN_CHAT_ID")
            return False
    return False


async def confirm_order(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    d = ctx.user_data
    user = update.effective_user
    if not SETTINGS["open"]:  # 客人填資料的途中剛好關店
        await update.message.reply_text(closed_text(), reply_markup=MAIN_KB)
        return ConversationHandler.END
    t = "d" if d["order_type"] == BTN_DELIVERY else "p"
    cart = get_cart(ctx)

    # 取號並寫入紀錄（加鎖：避免兩位客人同時下單拿到同一個編號，也避免同時搶最後一份庫存）
    problems = []
    async with ORDER_LOCK:
        problems = stock_problems(cart)
        if not problems:
            no = f"{next_order_no():04d}"
            now = datetime.now(TIMEZONE)
            log_event({
                "event": "order",
                "date": today_str(),
                "time": now.strftime("%H:%M:%S"),
                "ts": now.isoformat(),
                "session": SETTINGS["session"],
                "no": no,
                "uid": user.id,
                "name": d["name"],
                "phone": d["phone"],
                "type": t,
                "addr": d["addr"],
                "total": cart_total(cart),
                "rate": SETTINGS.get("rm_rate"),
                "rm_total": rm_amount(cart_total(cart), SETTINGS.get("rm_rate"), total=True),
                "lines": order_lines(cart),
            })
    if problems:  # 最後一刻被別人買走了
        await update.message.reply_text(stock_warning(problems), reply_markup=MAIN_KB)
        return ConversationHandler.END

    summary = (
        f"🧾 訂單 #{no}\n{cart_text(cart)}\n\n"
        f"方式：{d['order_type']}\n姓名：{d['name']}\n電話：{d['phone']}\n地址：{d['addr']}"
    )

    # 通知商家（附狀態按鈕）；失敗時告知客人，並把這筆訂單標記為作廢
    if ADMIN_CHAT_ID:
        ok = await send_to_admin(
            ctx.bot,
            f"🔔 新訂單\n\n{summary}\n\n👤 Telegram：{user.full_name} (id {user.id})",
            status_markup("new", user.id, no, t),
        )
        if not ok:
            log_event({"event": "status", "uid": user.id, "no": no, "status": "reject"})
            await update.message.reply_text(
                "⚠️ 目前無法送出訂單，請直接聯絡店家。\n" + SHOP_CONTACT,
                reply_markup=MAIN_KB,
            )
            return ConversationHandler.END

    # 記錄到 CSV
    try:
        save_order(no, d, cart)
    except Exception:
        logging.exception("寫入 orders.csv 失敗（若檔案正被 Excel 開啟，請先關閉）")

    # 這張單讓某個品項賣完了 → 通知商家
    if ADMIN_CHAT_ID:
        for bid in {base_id(parse_key(k)[0]) for k in cart}:
            if STOCK["limits"].get(bid) is not None and is_sold_out(bid):
                try:
                    await ctx.bot.send_message(
                        ADMIN_CHAT_ID,
                        f"📢 {ITEMS[bid][0]} 已售完（今日上限 {STOCK['limits'][bid]} 份），菜單已自動標示「已售完」。",
                    )
                except Exception:
                    logging.exception("傳送售完通知失敗")

    d.setdefault("orders", []).append(summary)
    cart.clear()
    await update.message.reply_text(
        f"✅ 訂單 #{no} 已送出，商家確認後會通知你，謝謝！\n\n"
        "📞 下單完成後主動聯繫客服核對訂單資訊",
        reply_markup=MAIN_KB,
    )
    if CS_USERNAME:
        await update.message.reply_text(
            "👇 點此聯繫客服",
            reply_markup=InlineKeyboardMarkup(
                [[InlineKeyboardButton("💬 聯繫客服", url=f"https://t.me/{CS_USERNAME}")]]
            ),
        )
    return ConversationHandler.END


async def cancel(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("已取消結帳，購物車內容保留。", reply_markup=MAIN_KB)
    return ConversationHandler.END


async def back_to_menu(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """結帳途中按「返回菜單」：離開結帳流程，回到菜單，購物車內容保留"""
    if await closed_guard(update):
        return ConversationHandler.END
    await update.message.reply_text("已返回菜單，購物車內容保留 👇", reply_markup=MAIN_KB)
    await update.message.reply_text("請選擇分類：", reply_markup=category_markup())
    return ConversationHandler.END


# ==================== 杯貼（NIIMBOT B21，50×30mm）====================
# 50×30mm、203 dpi → 400×240 點。版面以 400×240 設計，再乘上 STICKER_SCALE。
FONT_CANDIDATES = [
    r"C:\Windows\Fonts\msjhbd.ttc",  # 微軟正黑體 粗體
    r"C:\Windows\Fonts\msjh.ttc",    # 微軟正黑體
    r"C:\Windows\Fonts\kaiu.ttf",    # 標楷體
    r"C:\Windows\Fonts\mingliub.ttc",
    "/System/Library/Fonts/PingFang.ttc",
    "/System/Library/Fonts/STHeiti Medium.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/noto/NotoSansCJK-Bold.ttc",
    "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
]


def find_font():
    for path in ([STICKER_FONT] if STICKER_FONT else []) + FONT_CANDIDATES:
        if path and os.path.exists(path):
            return path
    return None


def sticker_items(order: dict) -> list:
    """把一張訂單展開成「一杯一張貼紙」"""
    cups = []
    for l in order.get("lines", []):
        if not (STICKER_ALL_ITEMS or is_drink(l["iid"])):
            continue
        cups.extend([l] * l["qty"])
    out = []
    for i, l in enumerate(cups, 1):
        parts = []
        if l.get("topping"):
            parts.append(TOPPING_NAME)
        if l.get("sugar"):
            parts.append(sugar_text(l["sugar"]))
        if l.get("ice"):
            parts.append(l["ice"])
        out.append({
            "no": order["no"],
            "name": l["name"],
            "topping": bool(l.get("topping")),
            "sugar": sugar_text(l.get("sugar", "")),
            "ice": l.get("ice", ""),
            "spec": "｜".join(parts),
            "customer": order.get("name", ""),
            "type": "外送" if order.get("type") == "d" else "外帶",
            "idx": i,
            "total": len(cups),
        })
    return out


def render_sticker(item: dict) -> bytes:
    """畫出一張 50×30mm 的杯貼，回傳 PNG 位元組"""
    from PIL import Image, ImageDraw, ImageFont

    font_path = find_font()
    if not font_path:
        raise RuntimeError("找不到中文字型，請在設定區填 STICKER_FONT（字型檔路徑）")
    k = max(1, int(STICKER_SCALE))
    W, H, M = 400 * k, 240 * k, 18 * k
    img = Image.new("L", (W, H), 255)
    d = ImageDraw.Draw(img)

    def font(size):
        return ImageFont.truetype(font_path, int(size * k))

    def fit(text, max_w, size, min_size):
        while size > min_size and d.textlength(text, font=font(size)) > max_w:
            size -= 2
        return font(size)

    def clip(text, max_w, f):
        if d.textlength(text, font=f) <= max_w:
            return text
        while text and d.textlength(text + "…", font=f) > max_w:
            text = text[:-1]
        return text + "…"

    # 第 1 行：單號（左）＋外帶/外送標籤（右，黑底白字）
    d.text((M, 34 * k), f"#{item['no']}", font=font(58), fill=0, anchor="lm")
    tag_f = font(34)
    tag_w = d.textlength(item["type"], font=tag_f) + 24 * k
    d.rounded_rectangle([W - M - tag_w, 10 * k, W - M, 58 * k], radius=8 * k, fill=0)
    d.text((W - M - tag_w / 2, 34 * k), item["type"], font=tag_f, fill=255, anchor="mm")

    # 第 2 行：品名（太長就縮小字）
    f = fit(item["name"], W - 2 * M, 50, 26)
    d.text((M, 98 * k), clip(item["name"], W - 2 * M, f), font=f, fill=0, anchor="lm")

    # 第 3 行：加料｜甜度｜冰塊
    if item["spec"]:
        f = fit(item["spec"], W - 2 * M, 38, 20)
        d.text((M, 152 * k), item["spec"], font=f, fill=0, anchor="lm")

    # 分隔線
    d.line([(M, 186 * k), (W - M, 186 * k)], fill=0, width=2 * k)

    # 第 4 行：客人姓名（左）＋第幾杯/共幾杯（右）
    idx_f = font(28)
    idx_text = f"{item['idx']}/{item['total']}"
    d.text((W - M, 208 * k), idx_text, font=idx_f, fill=0, anchor="rm")
    if STICKER_SHOW_NAME and item["customer"]:
        name_f = font(26)
        room = W - 2 * M - d.textlength(idx_text, font=idx_f) - 16 * k
        d.text((M, 208 * k), clip(item["customer"], room, name_f), font=name_f, fill=0, anchor="lm")

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def sticker_excel(items: list) -> bytes:
    """給 NIIMBOT App 批量列印用的 Excel：一杯一列"""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "貼紙"
    ws.append(["單號", "序號", "品名", "加料", "甜度", "冰塊", "規格", "姓名", "方式"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for it in items:
        ws.append([
            f"#{it['no']}", f"{it['idx']}/{it['total']}", it["name"],
            TOPPING_NAME if it["topping"] else "", it["sugar"], it["ice"],
            it["spec"], it["customer"] if STICKER_SHOW_NAME else "", it["type"],
        ])
    for col, w in zip("ABCDEFGHI", (9, 7, 22, 9, 9, 9, 22, 12, 7)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def send_retry(fn, *args, **kwargs):
    """傳送遇到 Telegram 限流時，等待後重試"""
    for _ in range(4):
        try:
            return await fn(*args, **kwargs)
        except RetryAfter as e:
            wait = e.retry_after
            wait = wait.total_seconds() if hasattr(wait, "total_seconds") else float(wait)
            await asyncio.sleep(wait + 1)
    return await fn(*args, **kwargs)


async def send_stickers(bot, chat_id: int, order: dict, note: str = ""):
    """把這張訂單的杯貼（圖片＋Excel）傳到商家群組。沒有飲料就什麼都不傳"""
    items = sticker_items(order)
    if not items:
        return
    no, n = order["no"], len(items)
    head = f"🏷 訂單 #{no} 杯貼 {n} 張{note}"
    try:
        if STICKER_SEND_IMAGES:
            images = await asyncio.to_thread(lambda: [render_sticker(it) for it in items])
            if n == 1:
                await send_retry(bot.send_photo, chat_id, photo=images[0], caption=head)
            else:
                for start in range(0, n, 10):  # 相簿一次最多 10 張
                    chunk = images[start:start + 10]
                    media = [
                        InputMediaPhoto(img, caption=head if (start == 0 and i == 0) else None)
                        for i, img in enumerate(chunk)
                    ]
                    await send_retry(bot.send_media_group, chat_id, media=media)
        if STICKER_SEND_EXCEL:
            data = await asyncio.to_thread(sticker_excel, items)
            await send_retry(
                bot.send_document, chat_id, document=io.BytesIO(data),
                filename=f"貼紙_{no}.xlsx",
                caption=f"📊 訂單 #{no} 批量列印檔（{n} 杯）",
            )
    except ImportError as e:
        await bot.send_message(chat_id, f"⚠️ 缺少套件，無法產生杯貼：{e.name}\n請在電腦執行：py -m pip install pillow openpyxl")
    except Exception as e:
        logging.exception("產生杯貼失敗")
        await bot.send_message(chat_id, f"⚠️ 訂單 #{no} 的杯貼產生失敗：{e}")


def find_order(uid, no):
    """依客人 id 與單號找到最近的一張訂單"""
    for o in reversed(load_all_orders()):
        if o.get("uid") == uid and o.get("no") == no:
            return o
    return None


async def reprint_stickers(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """「補印 12」或 /sticker 12：重傳目前場次中 #0012 的杯貼"""
    msg = update.effective_message
    m = re.search(r"(\d{1,4})\s*$", msg.text or "")
    if not m:
        await msg.reply_text("請寫單號，例如：補印 12")
        return
    no = f"{int(m.group(1)):04d}"
    orders = [o for o in load_orders_for(today_str()) if o.get("no") == no]
    if not orders:
        await msg.reply_text(f"目前場次（{today_str()}）找不到訂單 #{no}。")
        return
    order = orders[-1]
    if order.get("status") == "reject":
        await msg.reply_text(f"訂單 #{no} 已經拒單，不補印。")
        return
    if not sticker_items(order):
        await msg.reply_text(f"訂單 #{no} 沒有需要貼的飲料。")
        return
    await send_stickers(ctx.bot, msg.chat_id, order, note="（補印）")


async def sticker_test(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """「貼紙測試」：傳一張範例貼紙和範例 Excel，用來測試尺寸與列印效果（不會影響訂單）"""
    ids = list(MENU[DRINK_CATEGORY])[:2]  # 取菜單裡前兩個飲料當範例
    order = {
        "no": "0001", "name": "王小明", "type": "p",
        "lines": [
            {"iid": ids[0], "name": ITEMS[ids[0]][0], "topping": True, "sugar": SUGAR[2], "ice": ICE[2], "qty": 1, "price": 0},
            {"iid": ids[-1], "name": ITEMS[ids[-1]][0], "topping": False, "sugar": SUGAR[-1], "ice": ICE[-1], "qty": 1, "price": 0},
        ],
    }
    await send_stickers(ctx.bot, update.effective_message.chat_id, order, note="（測試）")


# ==================== 商家端：訂單狀態 ====================
STATUS_LABEL = {
    "accept": "✅ 已接單",
    "reject": "❌ 已拒單",
    "ready": "🍗 已出單",
    "done": "✔️ 已完成",
}


def status_markup(stage, uid, no, t):
    """依目前階段顯示下一步按鈕。callback 格式 st:動作:客人id:訂單號:類型(p外帶/d外送)"""
    def btn(text, action):
        return InlineKeyboardButton(text, callback_data=f"st:{action}:{uid}:{no}:{t}")

    if stage == "new":
        return InlineKeyboardMarkup([[btn("✅ 接單", "accept"), btn("❌ 拒單", "reject")]])
    if stage == "accept":
        label = "🛵 出單・通知客人外送中" if t == "d" else "🍗 出單・通知客人取餐"
        return InlineKeyboardMarkup([[btn(label, "ready")]])
    if stage == "ready":
        return InlineKeyboardMarkup([[btn("✔️ 已完成", "done")]])
    return None  # reject / done：不再顯示按鈕


def customer_message(action, no, t):
    if action == "accept":
        return f"✅ 訂單 #{no} 商家已接單，正在為你準備。"
    if action == "reject":
        return f"😢 抱歉，訂單 #{no} 目前無法接單，如有疑問請聯絡店家：\n{SHOP_CONTACT}"
    if action == "ready":
        if t == "d":
            return f"🛵 訂單 #{no} 已出單，外送人員已出發，請留意電話。"
        return f"🍗 訂單 #{no} 已出單，餐點熱騰騰，歡迎來取餐！"
    return f"✔️ 訂單 #{no} 已完成，謝謝光臨，歡迎再來！"


async def on_status(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    # 只有商家的對話（個人或群組）可以操作
    if ADMIN_CHAT_ID and q.message.chat_id != ADMIN_CHAT_ID:
        await q.answer("沒有權限", show_alert=True)
        return

    _, action, uid, no, t = q.data.split(":")
    uid = int(uid)
    label = STATUS_LABEL[action]

    base = q.message.text.split("\n\n📌")[0]
    stamp = datetime.now(TIMEZONE).strftime("%H:%M")
    new_text = f"{base}\n\n📌 狀態：{label}（{q.from_user.first_name} {stamp}）"

    await q.answer(label)
    await q.edit_message_text(new_text, reply_markup=status_markup(action, uid, no, t))

    # 記錄狀態（結算時，已拒單的訂單不計入）
    log_event({"event": "status", "uid": uid, "no": no, "status": action,
               "time": datetime.now(TIMEZONE).strftime("%H:%M:%S")})

    try:
        await ctx.bot.send_message(uid, customer_message(action, no, t))
    except Exception:
        await q.message.reply_text(
            f"⚠️ 無法通知訂單 #{no} 的客人（可能已封鎖機器人），請自行電話聯絡。"
        )

    # 接單 → 自動傳這張單的杯貼（背景執行，不卡住按鈕）
    if action == "accept" and STICKER_ENABLED:
        order = find_order(uid, no)
        if order:
            ctx.application.create_task(send_stickers(ctx.bot, q.message.chat_id, order))


# ==================== 結算報表 ====================
STATUS_TEXT = {
    "new": "待處理", "accept": "製作中", "ready": "已出單",
    "done": "已完成", "reject": "已拒單",
}


def order_rate(o: dict):
    """這張單當時的匯率；舊訂單沒有紀錄就用現在的匯率"""
    return o.get("rate") or SETTINGS.get("rm_rate")


def order_rm(o: dict):
    """這張單的馬幣總額（含 5 仙進位）；沒有匯率回傳 None"""
    if o.get("rm_total") is not None:
        return o["rm_total"]
    return rm_amount(o["total"], SETTINGS.get("rm_rate"), total=True)


def aggregate(day: str) -> dict:
    orders = load_orders_for(day)
    valid = [o for o in orders if o["status"] != "reject"]
    items, specs = {}, {}
    for o in valid:
        rate = order_rate(o)
        for l in o["lines"]:
            it = items.setdefault(
                l["iid"], {"name": l["name"], "qty": 0, "topping": 0, "amount": 0.0, "amount_rm": 0.0}
            )
            it["qty"] += l["qty"]
            it["amount"] += l["price"] * l["qty"]
            if rate:
                it["amount_rm"] += l["price"] * l["qty"] * rate
            if l["topping"]:
                it["topping"] += l["qty"]
            sk = (l["iid"], l["name"], bool(l["topping"]), l["sugar"], l["ice"])
            specs[sk] = specs.get(sk, 0) + l["qty"]
    done = sum(1 for o in valid if o["status"] == "done")
    rms = [order_rm(o) for o in valid]
    revenue_rm = None if (not valid or any(x is None for x in rms)) else round(sum(rms), 2)
    return {
        "day": day,
        "revenue_rm": revenue_rm,
        "orders": orders,
        "valid": valid,
        "items": items,
        "specs": specs,
        "done": done,
        "pending": len(valid) - done,
        "rejected": len(orders) - len(valid),
        "revenue": round(sum(o["total"] for o in valid), 2),
        "pickup": sum(1 for o in valid if o["type"] == "p"),
        "delivery": sum(1 for o in valid if o["type"] == "d"),
    }


def fmt_revenue(a: dict) -> str:
    base = f"{a['revenue']:.2f}{CURRENCY}"
    return base if a.get("revenue_rm") is None else f"{base}（RM{a['revenue_rm']:.2f}）"


def ordered_iids(items: dict) -> list:
    base = [iid for group in MENU.values() for iid in group]
    return base + [i for i in items if i not in base]


def line_label(l: dict) -> str:
    s = l["name"]
    if l["topping"]:
        s += f"({TOPPING_NAME})"
    if l["sugar"] or l["ice"]:
        s += f"（{'・'.join(p for p in (sugar_text(l['sugar']), l['ice']) if p)}）"
    return s


def summary_text(a: dict) -> str:
    if not a["orders"]:
        return f"📊 場次 {a['day']} 目前沒有訂單。"
    lines = [
        f"📊 結算｜場次 {a['day']}",
        "",
        f"🧾 訂單數：{len(a['valid'])}（已完成 {a['done']}・處理中 {a['pending']}）",
    ]
    if a["rejected"]:
        lines.append(f"❌ 已拒單：{a['rejected']}（不計入）")
    lines += [
        f"💰 營業額：{fmt_revenue(a)}",
        f"🏃 外帶 {a['pickup']}　🛵 外送 {a['delivery']}",
        "",
    ]
    total_qty = sum(i["qty"] for i in a["items"].values())
    lines.append(f"🍽 品項銷售（共 {total_qty} 份）")

    def row(iid):
        it = a["items"].get(iid)
        name = it["name"] if it else ITEMS[iid][0]
        qty = it["qty"] if it else 0
        extra = f"（含{TOPPING_NAME} {it['topping']}）" if it and it["topping"] else ""
        return f"• {name} x{qty}{extra}"

    for cat, group in MENU.items():
        lines.append(f"\n{cat}")
        lines += [row(iid) for iid in group]
    others = [iid for iid in a["items"] if iid not in ITEM_CAT]
    if others:
        lines.append("\n其他")
        lines += [row(iid) for iid in others]
    return "\n".join(lines)


def write_excel(a: dict) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor="FFD966")

    def style_header(ws):
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center")

    def width(s):
        return sum(2 if ord(ch) > 255 else 1 for ch in str(s))

    def autowidth(ws):
        for col in ws.columns:
            w = max((width(c.value) for c in col if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(w + 2, 60)

    wb = Workbook()

    # --- 今日結算（上方是摘要，下方是每個品項賣出的份數）---
    ws = wb.active
    ws.title = "今日結算"
    for k, v in [
        ("場次", a["day"]),
        ("訂單數（不含拒單）", len(a["valid"])),
        (f"營業額（{CURRENCY}）", a["revenue"]),
        *([("營業額（RM）", a["revenue_rm"])] if a.get("revenue_rm") is not None else []),
        ("已完成 / 處理中", f"{a['done']} / {a['pending']}"),
        ("已拒單（不計入）", a["rejected"]),
        ("外帶 / 外送", f"{a['pickup']} / {a['delivery']}"),
    ]:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append([])
    show_rm = a.get("revenue_rm") is not None
    ws.append(
        ["分類", "品項", "銷售份數", f"其中{TOPPING_NAME}", f"金額（{CURRENCY}）"]
        + (["金額（RM）"] if show_rm else [])
    )
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center")
    sum_qty = sum_amount = sum_rm = 0
    for iid in ordered_iids(a["items"]):
        it = a["items"].get(iid)
        name = it["name"] if it else ITEMS[iid][0]
        qty = it["qty"] if it else 0
        top = it["topping"] if it else 0
        amount = round(it["amount"], 2) if it else 0
        amount_rm = round(it["amount_rm"], 2) if it else 0
        ws.append([ITEM_CAT.get(iid, "其他"), name, qty, top, amount] + ([amount_rm] if show_rm else []))
        ws.cell(row=ws.max_row, column=5).number_format = "0.00"
        if show_rm:
            ws.cell(row=ws.max_row, column=6).number_format = "0.00"
        sum_qty += qty
        sum_amount += amount
        sum_rm += amount_rm
    ws.append(["合計", "", sum_qty, "", round(sum_amount, 2)] + ([round(sum_rm, 2)] if show_rm else []))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.cell(row=ws.max_row, column=5).number_format = "0.00"
    if show_rm:
        ws.cell(row=ws.max_row, column=6).number_format = "0.00"
    autowidth(ws)

    # --- 規格明細（甜度 / 冰塊 / 加珍珠）---
    ws = wb.create_sheet("規格明細")
    ws.append(["品項", TOPPING_NAME, "甜度", "冰塊", "數量"])
    order_index = {iid: n for n, iid in enumerate(ordered_iids(a["items"]))}
    for (iid, name, topping, sugar, ice), qty in sorted(
        a["specs"].items(), key=lambda kv: (order_index.get(kv[0][0], 999), kv[0][2], kv[0][3], kv[0][4])
    ):
        ws.append([name, "是" if topping else ("否" if (sugar or ice) else ""), sugar, ice, qty])
    style_header(ws)
    autowidth(ws)

    # --- 訂單明細 ---
    ws = wb.create_sheet("訂單明細")
    ws.append(["單號", "時間", "姓名", "電話", "方式", "地址", "內容", f"金額（{CURRENCY}）", "金額（RM）", "狀態"])
    for o in a["orders"]:
        ws.append([
            f"#{o['no']}", o.get("time", ""), o["name"], o["phone"],
            "外送" if o["type"] == "d" else "外帶", o["addr"],
            "；".join(f"{line_label(l)} x{l['qty']}" for l in o["lines"]),
            o["total"], order_rm(o) if order_rm(o) is not None else "",
            STATUS_TEXT.get(o["status"], o["status"]),
        ])
    style_header(ws)
    for r in ws.iter_rows(min_row=2, min_col=8, max_col=9):
        for c in r:
            c.number_format = "0.00"
    autowidth(ws)

    os.makedirs(REPORT_DIR, exist_ok=True)
    path = os.path.join(REPORT_DIR, f"結算報表_{a['day']}.xlsx")
    try:
        wb.save(path)
    except PermissionError:  # 檔案正被 Excel 開啟時，改存成另一個檔名
        path = os.path.join(
            REPORT_DIR, f"結算報表_{a['day']}_{datetime.now(TIMEZONE).strftime('%H%M%S')}.xlsx"
        )
        wb.save(path)
    return path


def write_orders_csv(a: dict) -> str:
    """某場次的訂單清單（一張單一列），Excel 可直接開啟，中文不亂碼"""
    os.makedirs(REPORT_DIR, exist_ok=True)
    rows = [["場次", "時間", "單號", "姓名", "電話", "方式", "地址", "內容", f"金額（{CURRENCY}）", "金額（RM）", "狀態"]]
    for o in a["orders"]:
        rows.append([
            a["day"], o.get("time", ""), f"#{o['no']}", o["name"], o["phone"],
            "外送" if o["type"] == "d" else "外帶", o["addr"],
            "；".join(f"{line_label(l)} x{l['qty']}" for l in o["lines"]),
            o["total"], order_rm(o) if order_rm(o) is not None else "",
            STATUS_TEXT.get(o["status"], o["status"]),
        ])
    rows.append([])
    rows.append([
        "合計（不含拒單）", "", "", "", "", "", "", "", a["revenue"],
        a["revenue_rm"] if a.get("revenue_rm") is not None else "", f"{len(a['valid'])} 單",
    ])

    def save(path):
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            csv.writer(f).writerows(rows)

    path = os.path.join(REPORT_DIR, f"訂單_{a['day']}.csv")
    try:
        save(path)
    except PermissionError:  # 檔案正被 Excel 開啟時，改存成另一個檔名
        path = os.path.join(
            REPORT_DIR, f"訂單_{a['day']}_{datetime.now(TIMEZONE).strftime('%H%M%S')}.csv"
        )
        save(path)
    return path


async def send_orders_csv(update: Update, a: dict):
    path = await asyncio.to_thread(write_orders_csv, a)
    with open(path, "rb") as f:
        await update.effective_chat.send_document(
            document=f,
            filename=os.path.basename(path),
            caption=f"🧾 場次 {a['day']} 訂單清單（共 {len(a['orders'])} 單）",
        )


async def do_orders(update: Update, day: str = None):
    day = day or today_str()
    if not valid_day_label(day):
        await update.effective_message.reply_text(
            "場次格式請用 2026-10-09 或 2026-10-09-2 這樣的寫法（輸入「營業日」可以查看有哪些場次）。"
        )
        return
    a = await asyncio.to_thread(aggregate, day)
    if not a["orders"]:
        await update.effective_message.reply_text(f"場次 {day} 目前沒有訂單。")
        return
    await send_orders_csv(update, a)


async def orders_text(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """商家群組輸入「訂單」或「訂單 2026-10-08」：只傳訂單清單"""
    await do_orders(update, ctx.matches[0].group(1) if ctx.matches else None)


async def orders_cmd(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    await do_orders(update, ctx.args[0] if ctx.args else None)


async def do_settle(update: Update, day: str = None):
    day = day or today_str()
    if not valid_day_label(day):
        await update.effective_message.reply_text(
            "場次格式請用 2026-10-09 或 2026-10-09-2 這樣的寫法（輸入「營業日」可以查看有哪些場次）。"
        )
        return

    a = await asyncio.to_thread(aggregate, day)
    await update.effective_message.reply_text(summary_text(a))
    if not a["orders"]:
        return
    await send_orders_csv(update, a)
    try:
        path = await asyncio.to_thread(write_excel, a)
    except ImportError:
        await update.effective_message.reply_text(
            '⚠️ 還沒安裝 Excel 套件，請在電腦執行：py -m pip install openpyxl\n安裝後再輸入一次「結算」。'
        )
        return
    except Exception as e:
        logging.exception("產生報表失敗")
        await update.effective_message.reply_text(f"❌ 產生 Excel 失敗：{e}")
        return
    with open(path, "rb") as f:
        await update.effective_chat.send_document(
            document=f, filename=os.path.basename(path), caption=f"📎 場次 {day} 結算報表"
        )


async def settle_text(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """商家群組輸入「結算」或「結算 2026-10-08」"""
    day = ctx.matches[0].group(1) if ctx.matches else None
    await do_settle(update, day)


async def settle_cmd(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """/settle 或 /settle 2026-10-08（機器人收不到一般文字時的備用方式）"""
    await do_settle(update, ctx.args[0] if ctx.args else None)


# ==================== 頻道廣告 ====================
async def send_ad(bot, ad):
    kb = InlineKeyboardMarkup(
        [[InlineKeyboardButton("🍽 立即點餐", url=f"https://t.me/{BOT_USERNAME}")]]
    )
    photo = ad.get("photo")
    photo_path = os.path.join(BASE_DIR, photo) if photo else ""
    if photo_path and os.path.exists(photo_path):
        with open(photo_path, "rb") as f:
            await bot.send_photo(CHANNEL_ID, photo=f, caption=ad["text"], reply_markup=kb)
    else:
        if photo:
            print("⚠️ 找不到廣告圖片：", photo_path, "（改發純文字）")
        await bot.send_message(CHANNEL_ID, ad["text"], reply_markup=kb)


async def post_ad(ctx: ContextTypes.DEFAULT_TYPE):
    await send_ad(ctx.bot, ctx.job.data)


async def testad(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """/testad 發全部廣告；/testad 1 只發第 1 篇；/testad 2 只發第 2 篇"""
    if ADMIN_CHAT_ID and update.effective_chat.id != ADMIN_CHAT_ID:
        return
    targets = ADS
    if ctx.args and ctx.args[0].isdigit() and 1 <= int(ctx.args[0]) <= len(ADS):
        targets = [ADS[int(ctx.args[0]) - 1]]
    try:
        for ad in targets:
            await send_ad(ctx.bot, ad)
        await update.message.reply_text(f"✅ 已發送 {len(targets)} 則廣告到頻道")
    except Exception as e:
        await update.message.reply_text(f"❌ 發送失敗：{e}")


# ==================== 診斷：營業日與訂單編號 ====================
def fmt_iso(ts) -> str:
    try:
        return datetime.fromisoformat(ts).astimezone(TIMEZONE).strftime("%m-%d %H:%M")
    except (TypeError, ValueError):
        return "?"


async def day_info(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """商家群組輸入「營業日」或 /day：查看營業狀態、目前場次與訂單編號"""
    now = datetime.now(TIMEZONE)
    day = today_str()
    todays = load_orders_for(day)
    nums = []
    for o in todays:
        try:
            nums.append(int(o["no"]))
        except (KeyError, ValueError):
            pass
    counts = {}
    for o in load_all_orders():
        counts[order_day(o)] = counts.get(order_day(o), 0) + 1

    lines = [
        f"🕒 機器人時間：{now.strftime('%Y-%m-%d %H:%M')}（{TIMEZONE.key}）",
        f"🏪 目前狀態：{'🟢 營業中' if SETTINGS['open'] else '🔴 已打烊'}"
        f"（外帶 {'✅' if SETTINGS['pickup'] else '❌'}｜外送 {'✅' if SETTINGS['delivery'] else '❌'}）",
    ]
    lines.append(
        f"💱 匯率：1{CURRENCY} = RM{SETTINGS['rm_rate']:g}" if SETTINGS.get("rm_rate") else "💱 匯率：尚未設定（只顯示 U）"
    )
    if SETTINGS["session"]:
        extra = f"，關店 {fmt_iso(SETTINGS['closed_at'])}" if not SETTINGS["open"] else ""
        lines.append(f"📅 目前場次：{day}（開店 {fmt_iso(SETTINGS['opened_at'])}{extra}）")
    else:
        lines.append(f"📅 還沒用過「開店」，暫以 {DAY_START_HOUR:02d}:00 起算的營業日：{day}")
    lines.append(f"🧾 這個場次訂單：{len(todays)} 張")
    if todays:
        last = max(todays, key=lambda o: o.get("ts") or o.get("time", ""))
        lines.append(f"　最後一張 #{last['no']}（下單時間 {last.get('time', '?')}）")
    lines.append(f"🔢 下一張訂單編號：#{(max(nums, default=0) + 1):04d}")
    lines.append("\n📚 紀錄中各場次的訂單數：")
    for d in sorted(counts)[-6:]:
        lines.append(f"　{d}：{counts[d]} 單")
    lines.append(f"\n版本：{VERSION}\n{run_info()}")
    await update.effective_message.reply_text("\n".join(lines))


# ==================== 商家：開店 / 關店 ====================
def new_session_label(now) -> str:
    """開店時建立新場次。名稱用開店日期；如果這個名稱已經有訂單（含舊的測試單），就加 -2、-3"""
    base = now.strftime("%Y-%m-%d")
    used = {order_day(o) for o in load_all_orders()}
    label, n = base, 1
    while label in used:
        n += 1
        label = f"{base}-{n}"
    return label


def status_line() -> str:
    return (
        f"外帶 {'✅ 開放' if SETTINGS['pickup'] else '❌ 暫停'}｜"
        f"外送 {'✅ 開放' if SETTINGS['delivery'] else '❌ 暫停'}"
    )


async def open_shop(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    msg = update.effective_message
    if SETTINGS["open"] and SETTINGS["session"]:
        await msg.reply_text(
            f"🟢 目前已經是營業中（場次 {SETTINGS['session']}，開店 {fmt_iso(SETTINGS['opened_at'])}）。\n"
            "要重新開始新場次，請先輸入「關店」。"
        )
        return
    now = datetime.now(TIMEZONE)
    SETTINGS.update(
        open=True, session=new_session_label(now), opened_at=now.isoformat(), closed_at=None
    )
    save_settings()
    await msg.reply_text(
        f"🟢 開店囉！\n場次：{SETTINGS['session']}\n訂單編號從 #0001 開始\n{status_line()}\n\n"
        f"{stock_hint()}\n\n"
        "客人現在可以下單。輸入「關店」停止接單，「結算」查看報表。"
    )


async def close_shop(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    msg = update.effective_message
    if not SETTINGS["open"]:
        await msg.reply_text("🔴 目前已經是打烊狀態。要開始營業請輸入「開店」。")
        return
    now = datetime.now(TIMEZONE)
    if not SETTINGS["session"]:  # 從來沒用過「開店」：把目前的營業日固定成這一場
        SETTINGS["session"] = business_day(now)
        SETTINGS["opened_at"] = None
    SETTINGS.update(open=False, closed_at=now.isoformat())
    save_settings()
    a = await asyncio.to_thread(aggregate, SETTINGS["session"])
    await msg.reply_text(
        f"🔴 已關店，客人不能再下單。\n場次：{SETTINGS['session']}\n"
        f"本場訂單 {len(a['valid'])} 張，營業額 {fmt_revenue(a)}\n\n"
        "客人點菜單、購物車、結帳，都會顯示「已休店」。\n"
        "輸入「結算」可以取得完整報表。"
    )
    try:  # 順便把資料備份傳到群組
        await send_backup(ctx.bot, msg.chat_id, "（關店）")
    except Exception:
        logging.exception("關店備份失敗")


async def shop_text(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """商家群組輸入「開店」或「關店」"""
    if ctx.matches[0].group(1) == "開店":
        await open_shop(update, ctx)
    else:
        await close_shop(update, ctx)


async def shop_cmd(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """備用指令：/open（開店）、/close（關店）"""
    cmd = update.effective_message.text.split()[0].lstrip("/").split("@")[0]
    await (open_shop if cmd == "open" else close_shop)(update, ctx)


# ==================== 商家：每日庫存 ====================
def base_items() -> dict:
    return {iid: name for group in MENU.values() for iid, (name, _) in group.items()}


def resolve_item(query: str):
    """用品名（可只寫一部分，例如「雞排」）找品項。回傳 (品項代號, 錯誤說明)"""
    q = query.strip().replace(" ", "")
    items = base_items()
    if q.lower() in items:
        return q.lower(), None
    exact = [i for i, n in items.items() if n.replace(" ", "") == q]
    if len(exact) == 1:
        return exact[0], None
    hits = [i for i, n in items.items() if q and q in n.replace(" ", "")]
    if len(hits) == 1:
        return hits[0], None
    if not hits:
        return None, f"❓ 找不到「{query.strip()}」，請寫菜單上的品名（可以只寫一部分）"
    return None, f"⚠️ 「{query.strip()}」符合多項（{'、'.join(items[i] for i in hits)}），請寫完整一點"


def stock_line(iid: str) -> str:
    sold = sold_counts(today_str()).get(iid, 0)
    limit = STOCK["limits"].get(iid)
    name = base_items()[iid]
    if iid in manual_soldout():
        return f"• {name}：❌ 已售完（手動標示，已賣 {sold}）"
    if limit is None:
        return f"• {name}：不限量" + (f"（已賣 {sold}）" if sold else "")
    r = max(limit - sold, 0)
    if r <= 0:
        return f"• {name}：❌ 已售完（{sold}/{limit}）"
    return f"• {name}：剩 {r}／{limit}（已賣 {sold}）"


def stock_overview() -> str:
    lines = [f"📦 今日庫存｜場次 {today_str()}"]
    for cat, group in MENU.items():
        lines.append(f"\n{cat}")
        lines += [stock_line(iid) for iid in group]
    return "\n".join(lines)


STOCK_HELP = (
    "\n\n📝 用法：\n"
    "設定今天可賣數量（可一次設定多項，一行一個）：\n"
    "設定庫存\n雞排 50\n雞腿 30\n塔香鹽酥雞 不限\n\n"
    "標示賣完：售完 雞排\n恢復供應：補貨 雞排\n全部清除：清除庫存"
)


async def show_stock(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """商家群組輸入「庫存」或 /stock"""
    await update.effective_message.reply_text(stock_overview() + STOCK_HELP)


def parse_stock_text(text: str):
    """把「雞排 50」「雞腿：30」「鹽酥雞 不限」這類內容拆成 [(品名, 數量或None)]"""
    pairs, bad = [], []
    for chunk in re.split(r"[\n,，、;；]+", text):
        chunk = chunk.strip()
        if not chunk:
            continue
        m = re.match(r"^(.+?)[\s:：=]*(\d+|不限量|不限|無限)$", chunk)
        if m:
            qty = None if not m.group(2).isdigit() else int(m.group(2))
            pairs.append((m.group(1).strip(), qty))
        else:
            bad.append(chunk)
    return pairs, bad


async def set_stock(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """商家群組輸入「設定庫存」+ 每行「品名 數量」（也可以 /setstock 雞排 50）"""
    msg = update.effective_message
    text = msg.text or ""
    body = re.sub(r"^\s*(設定庫存|/setstock(?:@\w+)?)", "", text, count=1)
    pairs, bad = parse_stock_text(body)
    if not pairs and not bad:
        await msg.reply_text(stock_overview() + STOCK_HELP)
        return
    out = []
    for query, qty in pairs:
        iid, err = resolve_item(query)
        if err:
            out.append(err)
            continue
        name = base_items()[iid]
        if qty is None:
            STOCK["limits"].pop(iid, None)
            out.append(f"✅ {name}：改為不限量")
        else:
            STOCK["limits"][iid] = qty
            out.append(f"✅ {name}：今天可賣 {qty} 份" + ("（等於已售完）" if qty == 0 else ""))
        m = STOCK["manual"]  # 重新設定數量 = 解除手動售完
        if m.get("label") == today_str() and iid in m["items"]:
            m["items"].remove(iid)
    out += [f"❓ 看不懂：「{b}」（請寫成「雞排 50」）" for b in bad]
    save_stock()
    await msg.reply_text("\n".join(out) + "\n\n" + stock_overview())


async def mark_soldout(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """「售完 雞排」或「已售完 雞排」"""
    msg = update.effective_message
    query = re.sub(r"^\s*(已售完|售完|/soldout(?:@\w+)?)", "", msg.text or "", count=1).strip()
    if not query:
        await msg.reply_text("請寫品名，例如：售完 雞排")
        return
    iid, err = resolve_item(query)
    if err:
        await msg.reply_text(err)
        return
    m = STOCK["manual"]
    if m.get("label") != today_str():
        m["label"], m["items"] = today_str(), []
    if iid not in m["items"]:
        m["items"].append(iid)
    save_stock()
    await msg.reply_text(f"❌ {base_items()[iid]} 已標示售完，客人的菜單會顯示「已售完」。\n要恢復請輸入：補貨 {query}")


async def restock_item(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """「補貨 雞排」：解除手動售完"""
    msg = update.effective_message
    query = re.sub(r"^\s*(補貨|/restock(?:@\w+)?)", "", msg.text or "", count=1).strip()
    if not query:
        await msg.reply_text("請寫品名，例如：補貨 雞排")
        return
    iid, err = resolve_item(query)
    if err:
        await msg.reply_text(err)
        return
    name = base_items()[iid]
    m = STOCK["manual"]
    if m.get("label") == today_str() and iid in m["items"]:
        m["items"].remove(iid)
        save_stock()
    r = remaining(iid)
    if r is not None and r <= 0:
        await msg.reply_text(
            f"⚠️ {name} 的今日數量已經賣完了，還是顯示已售完。\n"
            f"要繼續賣，請調高數量，例如：設定庫存 {name} 30"
        )
    else:
        await msg.reply_text(f"✅ {name} 恢復供應。\n" + stock_line(iid))


async def clear_stock(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """「清除庫存」：全部改回不限量"""
    STOCK["limits"].clear()
    STOCK["manual"] = {"label": None, "items": []}
    save_stock()
    await update.effective_message.reply_text("🧹 已清除所有庫存設定，全部品項改回不限量。")


def stock_hint() -> str:
    """開店時提醒目前的庫存設定"""
    limits = STOCK["limits"]
    if not limits:
        return "📦 目前沒有設定庫存（全部不限量），要限量請輸入「庫存」查看用法。"
    names = base_items()
    return "📦 今日庫存：" + "、".join(f"{names[i]} {n}" for i, n in limits.items() if i in names) + "\n（輸入「庫存」查看或修改）"


# ==================== 資料備份 / 匯入 ====================
BACKUP_FILES = ["orders_log.jsonl", "settings.json", "stock.json", "orders.csv"]


def make_backup():
    """把資料檔打包成 zip。回傳 (位元組, 有打包到的檔名)"""
    buf, names = io.BytesIO(), []
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name in BACKUP_FILES:
            path = os.path.join(DATA_DIR, name)
            if os.path.exists(path) and os.path.getsize(path) > 0:
                z.write(path, name)
                names.append(name)
    return buf.getvalue(), names


async def send_backup(bot, chat_id, note="") -> bool:
    data, names = await asyncio.to_thread(make_backup)
    if not names:
        return False
    stamp = datetime.now(TIMEZONE).strftime("%Y-%m-%d_%H%M")
    await send_retry(
        bot.send_document, chat_id, document=io.BytesIO(data), filename=f"備份_{stamp}.zip",
        caption=f"💾 資料備份{note}\n包含：{'、'.join(names)}\n要還原：把這個檔案傳到商家群組，說明文字寫「匯入」",
    )
    return True


async def backup_cmd(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """商家群組輸入「備份」或 /backup：傳資料備份檔"""
    ok = await send_backup(ctx.bot, update.effective_chat.id)
    if not ok:
        await update.effective_message.reply_text("目前還沒有可以備份的資料（還沒有訂單）。")


async def auto_backup(ctx: ContextTypes.DEFAULT_TYPE):
    try:
        await send_backup(ctx.bot, ADMIN_CHAT_ID, "（每日自動）")
    except Exception:
        logging.exception("自動備份失敗")


def write_data_file(name: str, data: bytes):
    path = os.path.join(DATA_DIR, name)
    tmp = path + ".tmp"
    with open(tmp, "wb") as f:
        f.write(data)
    os.replace(tmp, path)


async def import_backup(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """商家群組傳一個檔案、說明文字寫「匯入」：還原備份 zip，或單一資料檔"""
    msg = update.effective_message
    doc = msg.document
    name = doc.file_name or ""
    if (doc.file_size or 0) > 20 * 1024 * 1024:
        await msg.reply_text("⚠️ 檔案太大（超過 20MB），機器人無法下載。")
        return
    data = bytes(await (await doc.get_file()).download_as_bytearray())
    restored = []
    try:
        if name.lower().endswith(".zip"):
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                for info in z.infolist():
                    base = os.path.basename(info.filename)
                    if base in BACKUP_FILES and info.file_size <= 100 * 1024 * 1024:
                        write_data_file(base, z.read(info))
                        restored.append(base)
        elif name in BACKUP_FILES:
            write_data_file(name, data)
            restored.append(name)
    except zipfile.BadZipFile:
        await msg.reply_text("⚠️ 這個 zip 檔打不開，請用「備份」產生的檔案。")
        return
    if not restored:
        await msg.reply_text(
            "⚠️ 沒有可以匯入的檔案。請傳「備份」產生的 zip，或這些檔名之一：" + "、".join(BACKUP_FILES)
        )
        return
    SETTINGS.clear()
    SETTINGS.update(load_settings())
    STOCK.clear()
    STOCK.update(load_stock())
    _SOLD_CACHE.clear()
    await msg.reply_text(
        f"✅ 已匯入：{'、'.join(restored)}\n"
        f"目前場次：{today_str()}（{'營業中' if SETTINGS['open'] else '已打烊'}），這個場次有 {len(load_orders_for(today_str()))} 張訂單。\n"
        "輸入「營業日」可以確認。"
    )


# ==================== 商家：匯率（U → 馬幣）====================
async def set_rate(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """「匯率」查看；「匯率 4.45」設定 1U = RM4.45；「匯率 關閉」不顯示馬幣"""
    msg = update.effective_message
    m = re.match(r"^\s*(?:匯率|/rate(?:@\w+)?)\s*(.*)$", msg.text or "", re.S)
    arg = (m.group(1).strip() if m else "")
    cur = SETTINGS.get("rm_rate")
    if not arg:
        if cur:
            await msg.reply_text(f"💱 目前匯率：1{CURRENCY} = RM{cur:g}\n要修改請輸入：匯率 4.45　要關閉馬幣顯示：匯率 關閉")
        else:
            await msg.reply_text(
                f"💱 還沒設定匯率，客人目前只看到 {CURRENCY}。\n"
                f"設定方式：匯率 4.45（代表 1{CURRENCY} = RM4.45）"
            )
        return
    if arg in ("關閉", "關", "off", "0"):
        SETTINGS["rm_rate"] = None
        save_settings()
        await msg.reply_text(f"✅ 已關閉馬幣顯示，客人只會看到 {CURRENCY}。")
        return
    try:
        rate = float(arg.replace(",", "."))
    except ValueError:
        await msg.reply_text("格式不對，請寫成：匯率 4.45")
        return
    if not 0.1 <= rate <= 100:
        await msg.reply_text("這個數字看起來不太對（1U 應該是幾馬幣），請再確認，例如：匯率 4.45")
        return
    SETTINGS["rm_rate"] = rate
    save_settings()
    sample = list(base_items().items())[:2]
    lines = [f"✅ 已設定匯率：1{CURRENCY} = RM{rate:g}", ""]
    lines += [f"• {name}　{money(ITEMS[iid][1])}" for iid, name in sample]
    lines.append("\n客人現在看到的價格會同時顯示馬幣；之前的訂單仍依當時的匯率計算。")
    await msg.reply_text("\n".join(lines))


# ==================== 商家：開關外帶 / 外送 ====================
async def apply_toggle(update: Update, kind: str, enable: bool):
    SETTINGS["pickup" if kind == "外帶" else "delivery"] = enable
    save_settings()

    await update.effective_message.reply_text(
        f"{'✅ 已開放' if enable else '⏸ 已暫停'}{kind}\n\n目前狀態：{status_line()}"
    )


async def toggle_text(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """商家群組輸入「關閉外帶」「開啟外帶」「關閉外送」「開啟外送」"""
    m = ctx.matches[0]
    await apply_toggle(update, m.group(2), m.group(1) == "開啟")


async def toggle_cmd(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """備用指令：/pickup off、/pickup on、/delivery off、/delivery on"""
    cmd = update.effective_message.text.split()[0].lstrip("/").split("@")[0]
    kind = "外帶" if cmd == "pickup" else "外送"
    arg = ctx.args[0].lower() if ctx.args else ""
    if arg not in ("on", "off"):
        await update.effective_message.reply_text(f"用法：/{cmd} on（開放）或 /{cmd} off（暫停）")
        return
    await apply_toggle(update, kind, arg == "on")


# ==================== 速度診斷 ====================
def timed(fn):
    """處理超過 1.5 秒就在視窗印出提醒，方便找出哪裡慢"""
    async def wrapper(update, ctx):
        t0 = _time.monotonic()
        try:
            return await fn(update, ctx)
        finally:
            dt = _time.monotonic() - t0
            if dt > 1.5:
                print(f"⚠️ {fn.__name__} 處理了 {dt:.1f} 秒（偏慢，通常是網路延遲）")
    return wrapper


async def ping(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """商家群組輸入 /ping：測量訊息送達與機器人回覆的速度"""
    msg = update.effective_message
    lag = (datetime.now(timezone.utc) - msg.date).total_seconds()
    t0 = _time.monotonic()
    reply = await msg.reply_text("🏓 測試中…")
    rtt = _time.monotonic() - t0
    verdict = "✅ 速度正常" if lag < 3 and rtt < 3 else "⚠️ 偏慢，機器人所在的網路可能不穩"
    await reply.edit_text(
        f"🏓 訊息送達機器人：{max(lag, 0):.1f} 秒\n機器人回覆耗時：{rtt:.1f} 秒\n{verdict}\n\n{run_info()}"
    )


# ==================== 錯誤處理 ====================
def warn_once(key: str, text: str, every: int = 60):
    now = _time.monotonic()
    if now - ERROR_SEEN.get(key, -1e9) >= every:
        ERROR_SEEN[key] = now
        print(text)


async def on_error(update, ctx: ContextTypes.DEFAULT_TYPE):
    err = ctx.error
    if isinstance(err, Conflict):
        warn_once(
            "conflict",
            "⚠️⚠️ 偵測到另一份機器人正在同時運行（可能在別台電腦、VPS，或同事的電腦）！\n"
            "    同一個機器人只能開一份，否則訊息會被兩邊搶走，造成指令沒反應、訂單通知跑到別的地方。\n"
            "    請關掉其他的，只留一份。（在商家群組輸入 /ping 可以看是哪一台電腦在回覆）",
        )
        return
    if isinstance(err, BadRequest) and (
        "Query is too old" in str(err) or "not modified" in str(err)
    ):
        return  # 過期的按鈕點擊、重複點擊造成的「內容相同」，直接忽略
    if isinstance(err, NetworkError):  # 含 TimedOut：網路不穩，只印一行
        warn_once("network", f"⚠️ 網路不穩：{err}", every=10)
        return
    logging.error("發生錯誤", exc_info=err)


# ==================== 商家群組檢查 ====================
async def chatid(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """任何對話都能用：顯示這個對話的 chat id，並說明是不是機器人設定的商家群組"""
    chat = update.effective_chat
    ok = chat.id == ADMIN_CHAT_ID
    await update.effective_message.reply_text(
        f"💬 這個對話的 chat id：{chat.id}\n類型：{chat.type}\n\n"
        + (
            "✅ 這裡就是機器人設定的商家群組。"
            if ok
            else "⚠️ 這裡不是機器人設定的商家群組。\n"
            "如果這裡才是你的商家群組，請把 order_bot.py 的 ADMIN_CHAT_ID 改成上面這個數字，存檔後重新啟動。"
        )
    )


async def not_admin(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """商家指令在非商家群組被輸入時，說明原因（而不是靜悄悄地沒反應）"""
    chat = update.effective_chat
    await update.effective_message.reply_text(
        "⚠️ 這個指令只能在商家群組使用，而這裡不是機器人設定的商家群組。\n"
        f"這個對話的 chat id：{chat.id}\n\n"
        "如果這裡才是你的商家群組，請把 order_bot.py 的 ADMIN_CHAT_ID 改成上面這個數字，存檔後重新啟動。"
    )


async def on_migrate(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """群組升級成超級群組（chat id 改變）時，如果就是商家群組，自動改用新 id"""
    global ADMIN_CHAT_ID
    m = update.effective_message
    if m.migrate_to_chat_id:  # 舊群組收到的通知
        old, new = m.chat_id, m.migrate_to_chat_id
    elif m.migrate_from_chat_id:  # 新（超級）群組收到的通知
        old, new = m.migrate_from_chat_id, m.chat_id
    else:
        return
    if old == ADMIN_CHAT_ID:
        ADMIN_CHAT_ID = new
        if ADMIN_FILTER is not None and hasattr(ADMIN_FILTER, "add_chat_ids"):
            ADMIN_FILTER.add_chat_ids(new)
            ADMIN_FILTER.remove_chat_ids(old)
        SETTINGS.update(admin_migrated_from=old, admin_migrated_to=new)
        save_settings()
        print(f"🔄 商家群組已升級，chat id 由 {old} 變成 {new}，已自動更新。")
    elif new != ADMIN_CHAT_ID:  # 跟商家群組無關的群組升級
        return
    if m.migrate_from_chat_id:  # 只在新群組回覆
        await m.reply_text(
            f"🔄 群組升級了，機器人已自動改用新的 chat id：{new}\n"
            f"建議把 order_bot.py 的 ADMIN_CHAT_ID 也改成 {new}（重啟後仍然有效）。"
        )


# ==================== 主程式 ====================
def main():
    global ADMIN_FILTER
    print("機器人啟動｜版本：", VERSION)
    print(run_info())
    app = (
        Application.builder()
        .token(TOKEN)
        .concurrent_updates(True)
        .connect_timeout(15)
        .read_timeout(15)
        .write_timeout(15)
        .pool_timeout(15)
        .get_updates_connect_timeout(15)
        .get_updates_read_timeout(15)
        .build()
    )

    cancel_filter = filters.Regex(f"^{BTN_CANCEL}$") | filters.Command("cancel")
    text_filter = (
        filters.TEXT & ~filters.COMMAND
        & ~filters.Regex(f"^{BTN_CANCEL}$") & ~filters.Regex(f"^{BTN_BACK_MENU}$")
    )
    admin_filter = filters.Chat(chat_id=ADMIN_CHAT_ID) if ADMIN_CHAT_ID else filters.ALL
    ADMIN_FILTER = admin_filter

    checkout = ConversationHandler(
        entry_points=[
            MessageHandler(filters.Regex(f"^{BTN_CHECKOUT}$"), checkout_start),
            CallbackQueryHandler(checkout_start, pattern="^checkout$"),
        ],
        states={
            TYPE: [MessageHandler(filters.Regex(f"^({BTN_PICKUP}|{BTN_DELIVERY})$"), got_type)],
            NAME: [MessageHandler(text_filter, got_name)],
            PHONE: [MessageHandler(text_filter, got_phone)],
            ADDR: [MessageHandler(text_filter, got_addr)],
            CONFIRM: [MessageHandler(filters.Regex(f"^{BTN_CONFIRM}$"), timed(confirm_order))],
        },
        fallbacks=[
            MessageHandler(cancel_filter, cancel),
            MessageHandler(filters.Regex(f"^{BTN_BACK_MENU}$"), back_to_menu),
        ],
    )

    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("chatid", chatid))
    app.add_handler(MessageHandler(filters.StatusUpdate.MIGRATE, on_migrate))
    app.add_handler(CommandHandler("testad", testad, filters=admin_filter))
    # 結算：只有商家對話可以使用
    app.add_handler(CommandHandler("settle", settle_cmd, filters=admin_filter))
    app.add_handler(MessageHandler(
        filters.Regex(r"^\s*結算(?:\s+(\d{4}-\d{2}-\d{2}(?:-\d+)?))?\s*$") & admin_filter, settle_text
    ))
    app.add_handler(CommandHandler("sticker", reprint_stickers, filters=admin_filter))
    app.add_handler(MessageHandler(filters.Regex(r"^\s*(補印|貼紙)\s+#?\d{1,4}\s*$") & admin_filter, reprint_stickers))
    app.add_handler(CommandHandler("stickertest", sticker_test, filters=admin_filter))
    app.add_handler(MessageHandler(filters.Regex(r"^\s*貼紙測試\s*$") & admin_filter, sticker_test))
    app.add_handler(CommandHandler("backup", backup_cmd, filters=admin_filter))
    app.add_handler(MessageHandler(filters.Regex(r"^\s*備份\s*$") & admin_filter, backup_cmd))
    app.add_handler(MessageHandler(
        filters.Document.ALL & filters.CaptionRegex(r"^\s*匯入") & admin_filter, import_backup
    ))
    app.add_handler(CommandHandler("rate", set_rate, filters=admin_filter))
    app.add_handler(MessageHandler(filters.Regex(r"^\s*匯率(\s[\s\S]*)?$") & admin_filter, set_rate))
    app.add_handler(CommandHandler("hidemenu", hide_menu, filters=admin_filter))
    app.add_handler(MessageHandler(
        filters.Regex(r"^\s*(隱藏選單|關閉選單|收起選單)\s*$") & admin_filter, hide_menu
    ))
    # 每日庫存
    app.add_handler(CommandHandler("stock", show_stock, filters=admin_filter))
    app.add_handler(MessageHandler(filters.Regex(r"^\s*庫存\s*$") & admin_filter, show_stock))
    app.add_handler(CommandHandler("setstock", set_stock, filters=admin_filter))
    app.add_handler(MessageHandler(filters.Regex(r"^\s*設定庫存") & admin_filter, set_stock))
    app.add_handler(CommandHandler("soldout", mark_soldout, filters=admin_filter))
    app.add_handler(MessageHandler(filters.Regex(r"^\s*(已售完|售完)(\s|$)") & admin_filter, mark_soldout))
    app.add_handler(CommandHandler("restock", restock_item, filters=admin_filter))
    app.add_handler(MessageHandler(filters.Regex(r"^\s*補貨(\s|$)") & admin_filter, restock_item))
    app.add_handler(MessageHandler(filters.Regex(r"^\s*清除庫存\s*$") & admin_filter, clear_stock))
    app.add_handler(CommandHandler("ping", ping, filters=admin_filter))
    app.add_handler(CommandHandler("day", day_info, filters=admin_filter))
    app.add_handler(MessageHandler(
        filters.Regex(r"^\s*(開店|關店)\s*$") & admin_filter, shop_text
    ))
    app.add_handler(CommandHandler(["open", "close"], shop_cmd, filters=admin_filter))
    app.add_handler(MessageHandler(filters.Regex(r"^\s*營業日\s*$") & admin_filter, day_info))
    app.add_handler(MessageHandler(
        filters.Regex(r"^\s*(開啟|關閉)(外帶|外送)\s*$") & admin_filter, toggle_text
    ))
    app.add_handler(CommandHandler(["pickup", "delivery"], toggle_cmd, filters=admin_filter))
    app.add_handler(CommandHandler("orders", orders_cmd, filters=admin_filter))
    app.add_handler(MessageHandler(
        filters.Regex(r"^\s*訂單(?:\s+(\d{4}-\d{2}-\d{2}(?:-\d+)?))?\s*$") & admin_filter, orders_text
    ))
    # 商家指令在非商家群組被輸入時，說明原因與 chat id
    app.add_handler(CommandHandler(
        ["ping", "day", "settle", "orders", "open", "close", "pickup", "delivery", "testad",
         "stock", "setstock", "soldout", "restock", "hidemenu", "sticker", "stickertest", "rate", "backup"],
        not_admin, filters=~admin_filter,
    ))
    app.add_handler(MessageHandler(
        filters.Regex(r"(?s)^\s*(結算|訂單|營業日|開店|關店|庫存|設定庫存|已?售完|補貨|清除庫存|隱藏選單|關閉選單|收起選單|補印|貼紙測試|匯率|備份|(開啟|關閉)(外帶|外送))(\s.*)?$")
        & filters.ChatType.GROUPS & ~admin_filter,
        not_admin,
    ))
    app.add_handler(checkout)
    app.add_handler(MessageHandler(filters.Regex(f"^{BTN_MENU}$"), show_menu))
    app.add_handler(MessageHandler(filters.Regex(f"^{BTN_CART}$"), show_cart))
    app.add_handler(MessageHandler(filters.Regex(f"^{BTN_ORDERS}$"), show_orders))
    app.add_handler(MessageHandler(filters.Regex(f"^{BTN_CONTACT}$"), show_contact))
    app.add_handler(CallbackQueryHandler(timed(on_status), pattern="^st:"))  # 必須放在 on_callback 前面
    app.add_handler(CallbackQueryHandler(timed(on_callback)))
    app.add_error_handler(on_error)

    # 頻道定時廣告（每篇各自的發送時間）
    if CHANNEL_ID:
        if app.job_queue is None:
            print('⚠️ 尚未安裝排程功能，請執行：py -m pip install "python-telegram-bot[job-queue]"')
        else:
            for ad in ADS:
                for t in ad["times"]:
                    app.job_queue.run_daily(post_ad, time=t.replace(tzinfo=TIMEZONE), data=ad)
            if TEST_AD_ON_START:
                app.job_queue.run_once(post_ad, 10, data=ADS[0])

    if AUTO_BACKUP_TIME and ADMIN_CHAT_ID and app.job_queue is not None:
        app.job_queue.run_daily(auto_backup, time=AUTO_BACKUP_TIME.replace(tzinfo=TIMEZONE))

    app.run_polling(drop_pending_updates=True, bootstrap_retries=-1)


if __name__ == "__main__":
    main()
